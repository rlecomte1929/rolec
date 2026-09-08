from __future__ import annotations

import io
import os
import sqlite3
import tempfile
import unittest
from typing import Any, Dict, List, Optional

from fastapi import FastAPI
from fastapi.testclient import TestClient

try:
    from alembic import command
    from alembic.config import Config
except ImportError:  # pragma: no cover
    command = None
    Config = None

from backend.app.db import Base
from backend.app.models import (
    AssignmentType,
    BenefitCategory,
    Frequency,
    Phase,
    ProviderEntity,
    ValueType,
)
from backend.app.routers import policy_canonical as policy_canonical_router
from backend.app.services.policy_canonical_chunking import (
    build_canonical_chunks_from_elements,
    canonical_chunks_to_assistant_chunks,
)
from backend.app.services.policy_canonical_extraction import (
    OpenAIPolicyCanonicalExtractor,
    extract_canonical_policy_facts,
    project_canonical_fact_to_legacy_fact,
)
from backend.app.services.policy_canonical_ingestion import ingest_canonical_policy_document


def _minimal_docx_bytes() -> bytes:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise unittest.SkipTest("python-docx not installed") from exc

    doc = Document()
    doc.add_paragraph("Long Term Assignment Policy Summary")
    doc.add_paragraph("Housing allowance up to USD 5000 per month for long-term assignments.")
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Benefit"
    table.rows[0].cells[1].text = "Mobility premium 10%"
    table.rows[0].cells[2].text = "Section 2.1"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class _RecordingDb:
    def __init__(self) -> None:
        self.documents: Dict[str, Dict[str, Any]] = {}
        self.chunks: Dict[str, List[Dict[str, Any]]] = {}
        self.facts: List[Dict[str, Any]] = []
        self.errors: List[Dict[str, Any]] = []

    def insert_canonical_policy_document(self, **kwargs: Any) -> str:
        doc_id = "canon-doc-1"
        self.documents[doc_id] = {"id": doc_id, **kwargs}
        return doc_id

    def get_canonical_policy_document(self, document_id: str) -> Optional[Dict[str, Any]]:
        return self.documents.get(document_id)

    def delete_canonical_policy_artifacts(self, document_id: str) -> None:
        self.chunks[document_id] = []
        self.facts = [fact for fact in self.facts if fact["canonical_policy_document_id"] != document_id]
        self.errors = [err for err in self.errors if err["canonical_policy_document_id"] != document_id]

    def insert_canonical_policy_document_chunk(self, **kwargs: Any) -> str:
        chunk_id = f"chunk-{kwargs['chunk_index']}"
        row = {"id": chunk_id, **kwargs}
        self.chunks.setdefault(kwargs["canonical_policy_document_id"], []).append(row)
        return chunk_id

    def list_canonical_policy_document_chunks(self, document_id: str) -> List[Dict[str, Any]]:
        return self.chunks.get(document_id, [])

    def update_canonical_policy_document(self, document_id: str, **kwargs: Any) -> None:
        self.documents[document_id].update(kwargs)

    def insert_canonical_policy_fact(self, **kwargs: Any) -> str:
        fact_id = f"fact-{len(self.facts) + 1}"
        self.facts.append({"id": fact_id, **kwargs})
        return fact_id

    def insert_canonical_policy_validation_error(self, **kwargs: Any) -> str:
        error_id = f"error-{len(self.errors) + 1}"
        self.errors.append({"id": error_id, **kwargs})
        return error_id

    def list_canonical_policy_facts(self, document_id: str, **kwargs: Any) -> List[Dict[str, Any]]:
        _ = kwargs
        return [fact for fact in self.facts if fact["canonical_policy_document_id"] == document_id]

    def list_canonical_policy_validation_errors(self, document_id: str) -> List[Dict[str, Any]]:
        return [err for err in self.errors if err["canonical_policy_document_id"] == document_id]

    def get_canonical_policy_audit_summary(self, document_id: str) -> Dict[str, Any]:
        return {
            "document_id": document_id,
            "chunks_count": len(self.list_canonical_policy_document_chunks(document_id)),
            "facts_count": len(self.list_canonical_policy_facts(document_id)),
            "validation_error_count": len(self.list_canonical_policy_validation_errors(document_id)),
            "validation_pass_rate": 1.0,
            "counts_by_category": {},
            "counts_by_phase": {},
        }

    def get_active_canonical_policy_document_for_company(self, company_id: str) -> Optional[Dict[str, Any]]:
        for document in self.documents.values():
            if document.get("company_id") == company_id:
                return document
        return None

    def insert_canonical_policy_query_audit_log(self, **kwargs: Any) -> str:
        return "audit-1"


class _MockCompletionResponse:
    def __init__(self, content: str) -> None:
        self.choices = [type("Choice", (), {"message": type("Message", (), {"content": content})()})()]


class _MockOpenAIClient:
    def __init__(self, content: str) -> None:
        self.chat = type(
            "Chat",
            (),
            {
                "completions": type(
                    "Completions",
                    (),
                    {"create": lambda _self, **kwargs: _MockCompletionResponse(content)},
                )()
            },
        )()


class PolicyCanonicalPipelineTests(unittest.TestCase):
    def test_models_and_enums_registered(self) -> None:
        self.assertIn("canonical_policy_documents", Base.metadata.tables)
        self.assertIn("canonical_policy_document_chunks", Base.metadata.tables)
        self.assertIn("canonical_policy_facts", Base.metadata.tables)
        self.assertEqual(AssignmentType.LONG_TERM.value, "long_term")
        self.assertEqual(Phase.PRE_ASSIGNMENT.value, "pre_assignment")
        self.assertEqual(BenefitCategory.HOUSING.value, "housing")
        self.assertEqual(ValueType.MONETARY.value, "monetary")
        self.assertEqual(Frequency.MONTHLY.value, "monthly")
        self.assertEqual(ProviderEntity.COMPANY.value, "company")

    def test_alembic_migration_smoke(self) -> None:
        if command is None or Config is None:
            raise unittest.SkipTest("alembic not installed")
        # alembic/env.py prefers DATABASE_URL over the config's sqlalchemy.url.
        # In CI that env var points at ci_test.db — which `db.init_db()` has
        # already populated in the Prime step — so without clearing it here
        # upgrade() would target that pre-populated DB and collide on
        # CREATE TABLE canonical_policy_documents. Drop it for this call only.
        previous_db_url = os.environ.pop("DATABASE_URL", None)
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                db_path = os.path.join(tmpdir, "canonical.sqlite")
                seed_conn = sqlite3.connect(db_path)
                seed_conn.execute("CREATE TABLE policy_documents (id TEXT PRIMARY KEY)")
                seed_conn.commit()
                seed_conn.close()
                cfg = Config(os.path.join(os.path.dirname(__file__), "..", "alembic.ini"))
                cfg.set_main_option("script_location", os.path.join(os.path.dirname(__file__), "..", "alembic"))
                cfg.set_main_option("sqlalchemy.url", f"sqlite:///{db_path}")
                command.upgrade(cfg, "head")
                conn = sqlite3.connect(db_path)
                tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
                conn.close()
                self.assertIn("canonical_policy_documents", tables)
                self.assertIn("canonical_policy_document_chunks", tables)
                self.assertIn("canonical_policy_facts", tables)
        finally:
            if previous_db_url is not None:
                os.environ["DATABASE_URL"] = previous_db_url

    def test_ingestion_persists_canonical_document(self) -> None:
        db = _RecordingDb()
        document = ingest_canonical_policy_document(
            db,
            company_id="company-a",
            file_bytes=_minimal_docx_bytes(),
            filename="sample.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertEqual(document["company_id"], "company-a")
        self.assertEqual(document["filename"], "sample.docx")
        # Classifier output is structural (DocTypeProfile). The sample docx
        # begins with "Long Term Assignment Policy Summary" — an assignment
        # policy, not a generic policy summary. Prior expectation of
        # "policy_summary" encoded the old keyword-cascade behavior, removed
        # as part of the EXTRACT-BUG-2 fix.
        self.assertEqual(document["document_type"], "assignment_policy")
        self.assertIn("Housing allowance", document["normalized_text"])

    def test_chunk_builder_preserves_structure(self) -> None:
        elements = [
            {"text": "1. Benefits", "page": 1, "is_table_row": False},
            {"text": "Housing allowance up to USD 5000 per month.", "page": 1, "is_table_row": False},
            {"text": "Mobility premium | 10% | STA/LTA", "page": 1, "is_table_row": True},
        ]
        chunks = build_canonical_chunks_from_elements(elements)
        self.assertGreaterEqual(len(chunks), 1)
        self.assertTrue(any(chunk.get("structure_type") in {"paragraph", "mixed", "table_row"} for chunk in chunks))
        assistant_chunks = canonical_chunks_to_assistant_chunks(
            [{"id": "chunk-1", **chunk} for chunk in chunks]
        )
        self.assertEqual(assistant_chunks[0]["chunk_index"], 0)

    def test_extraction_with_mock_client_and_legacy_projection(self) -> None:
        db = _RecordingDb()
        db.documents["canon-doc-1"] = {
            "id": "canon-doc-1",
            "company_id": "company-a",
            "source_policy_document_id": None,
            "default_currency": "USD",
            "assignment_types_json": ["long_term"],
        }
        db.chunks["canon-doc-1"] = [
            {
                "id": "chunk-1",
                "company_id": "company-a",
                "chunk_index": 0,
                "section_path": "Benefits",
                "structure_type": "paragraph",
                "page_number": 1,
                "text_content": "Housing allowance up to USD 5000 per month.",
            }
        ]
        content = """
        {
          "facts": [
            {
              "benefit_category": "housing",
              "value_type": "monetary",
              "frequency": "monthly",
              "provider_entity": "company",
              "title": "Housing allowance",
              "description": "Monthly housing allowance",
              "assignment_types": ["long_term"],
              "amount": 5000,
              "currency": "USD",
              "source_quote": "Housing allowance up to USD 5000 per month.",
              "confidence_score": 0.91,
              "raw_payload": {"source": "mock"}
            }
          ]
        }
        """
        extractor = OpenAIPolicyCanonicalExtractor(client=_MockOpenAIClient(content), model="mock-model")
        result = extract_canonical_policy_facts(db, "canon-doc-1", extractor=extractor)
        self.assertEqual(result["facts_inserted"], 1)
        self.assertEqual(len(db.facts), 1)
        legacy = project_canonical_fact_to_legacy_fact(
            {
                "canonical_policy_document_chunk_id": "chunk-1",
                "benefit_category": "housing",
                "phase": None,
                "amount": 5000,
                "currency": "USD",
                "frequency": "monthly",
                "eligibility": {},
                "assignment_types": ["long_term"],
                "confidence_score": 0.91,
                "source_quote": "Housing allowance up to USD 5000 per month.",
            }
        )
        self.assertEqual(legacy["category"], "housing")

    def test_invalid_fact_is_logged_not_inserted(self) -> None:
        db = _RecordingDb()
        db.documents["canon-doc-1"] = {
            "id": "canon-doc-1",
            "company_id": "company-a",
            "source_policy_document_id": None,
            "default_currency": "USD",
            "assignment_types_json": [],
        }
        db.chunks["canon-doc-1"] = [
            {
                "id": "chunk-1",
                "company_id": "company-a",
                "chunk_index": 0,
                "section_path": "Benefits",
                "structure_type": "paragraph",
                "page_number": 1,
                "text_content": "Mobility premium 10%",
            }
        ]
        content = """
        {
          "facts": [
            {
              "benefit_category": "mobility_premium",
              "value_type": "percentage",
              "title": "Mobility premium",
              "description": "Percentage premium"
            }
          ]
        }
        """
        extractor = OpenAIPolicyCanonicalExtractor(client=_MockOpenAIClient(content), model="mock-model")
        result = extract_canonical_policy_facts(db, "canon-doc-1", extractor=extractor)
        self.assertEqual(result["facts_inserted"], 0)
        self.assertEqual(len(db.errors), 1)

    def test_router_lists_documents(self) -> None:
        original_get_user_by_token = policy_canonical_router.db.get_user_by_token
        original_get_profile_record = policy_canonical_router.db.get_profile_record
        original_is_admin_allowlisted = policy_canonical_router.db.is_admin_allowlisted
        original_list_docs = policy_canonical_router.db.list_canonical_policy_documents
        try:
            policy_canonical_router.db.get_user_by_token = lambda token: {"id": "u1", "role": "HR", "email": "hr@relopass.com"}
            policy_canonical_router.db.get_profile_record = lambda user_id: {"id": user_id, "company_id": "company-a", "role": "HR"}
            policy_canonical_router.db.is_admin_allowlisted = lambda email: True
            policy_canonical_router.db.list_canonical_policy_documents = lambda **kwargs: [{"id": "canon-doc-1"}]
            app = FastAPI()
            app.include_router(policy_canonical_router.admin_router, prefix="/api/admin")
            client = TestClient(app)
            response = client.get(
                "/api/admin/policy-canonical/documents?company_id=company-a",
                headers={"Authorization": "Bearer token"},
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()[0]["id"], "canon-doc-1")
        finally:
            policy_canonical_router.db.get_user_by_token = original_get_user_by_token
            policy_canonical_router.db.get_profile_record = original_get_profile_record
            policy_canonical_router.db.is_admin_allowlisted = original_is_admin_allowlisted
            policy_canonical_router.db.list_canonical_policy_documents = original_list_docs


if __name__ == "__main__":
    unittest.main()
